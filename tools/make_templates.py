#!/usr/bin/env python3
"""生成内置的空白 Office 模板（blank.docx / blank.xlsx / blank.pptx）。

依赖：pip install python-docx openpyxl python-pptx
用法：python3 tools/make_templates.py
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent.parent / "Resources"
EPOCH = datetime(2026, 1, 1)


def clean_core(props):
    """清空作者等元数据，避免模板携带生成工具或个人信息。"""
    for attr in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
        if hasattr(props, attr):
            setattr(props, attr, "")
    for attr in ("created", "modified"):
        if hasattr(props, attr):
            setattr(props, attr, EPOCH)


def make_docx():
    doc = Document()
    clean_core(doc.core_properties)
    doc.save(OUT / "blank.docx")


def make_xlsx():
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.created = EPOCH
    wb.properties.modified = EPOCH
    wb.save(OUT / "blank.xlsx")


def make_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])  # 一页空白幻灯片
    clean_core(prs.core_properties)
    prs.save(OUT / "blank.pptx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_docx()
    make_xlsx()
    make_pptx()
    for f in sorted(OUT.glob("blank.*")):
        print(f"{f.name}: {f.stat().st_size} bytes")
